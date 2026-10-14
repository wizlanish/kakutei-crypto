from pathlib import Path

from openpyxl import load_workbook

from kakutei_crypto.cli import main

FIX = Path(__file__).parent / "fixtures"


def import_fixtures(tmp_path, *names):
    ledger = tmp_path / "ledger.csv"
    assert main(["import", *(str(FIX / n) for n in names), "-o", str(ledger)]) == 0
    return ledger


def test_import_reports_unknown_files_and_still_writes_known_ones(tmp_path, capsys):
    mystery = tmp_path / "mystery.csv"
    mystery.write_text("a,b\n1,2\n", encoding="utf-8")
    ledger = tmp_path / "ledger.csv"
    code = main(["import", str(FIX / "okj_orders.csv"), str(mystery), "-o", str(ledger)])
    assert code == 2
    assert "mystery.csv" in capsys.readouterr().err
    assert ledger.read_text(encoding="utf-8").count("\n") == 3  # ヘッダー + 2 件


def test_check_fails_when_fee_is_paid_from_coins_never_acquired(tmp_path, capsys):
    ledger = import_fixtures(tmp_path, "okj_transfers.csv")
    assert main(["check", str(ledger)]) == 1
    assert "HBAR" in capsys.readouterr().out


def test_calc_refuses_to_compute_a_ledger_with_errors(tmp_path, capsys):
    ledger = import_fixtures(tmp_path, "okj_transfers.csv")
    assert main(["calc", str(ledger)]) == 1
    assert "所得金額" not in capsys.readouterr().out


def test_calc_prints_income_and_writes_xlsx(tmp_path, capsys):
    ledger = import_fixtures(tmp_path, "okj_orders.csv", "coincheck_youshiki1.csv")
    xlsx = tmp_path / "out.xlsx"
    assert main(["calc", str(ledger), "--xlsx", str(xlsx)]) == 0
    out = capsys.readouterr().out
    # BTC: 収入 661,000 − 原価 600,000 ／ ETH: 原価 600（手数料分）／ 報酬 300 ／ 円手数料 10
    assert "所得金額" in out and "60,690" in out
    book = load_workbook(xlsx)
    assert {"計算書_2026", "明細"} <= set(book.sheetnames)


def test_calc_accepts_multiple_ledgers_and_moving_method(tmp_path, capsys):
    a = import_fixtures(tmp_path, "okj_orders.csv")
    b = tmp_path / "b.csv"
    assert main(["import", str(FIX / "coincheck_youshiki1.csv"), "-o", str(b)]) == 0
    assert main(["calc", str(a), str(b), "--method", "moving", "--year", "2026"]) == 0
    assert "移動平均法" in capsys.readouterr().out


def test_import_refuses_to_overwrite_existing_ledger_without_force(tmp_path, capsys):
    ledger = import_fixtures(tmp_path, "okj_orders.csv")
    ledger.write_text(ledger.read_text(encoding="utf-8") + "", encoding="utf-8")
    before = ledger.read_text(encoding="utf-8")
    assert main(["import", str(FIX / "coincheck_youshiki1.csv"), "-o", str(ledger)]) == 1
    assert ledger.read_text(encoding="utf-8") == before
    assert "--force" in capsys.readouterr().err
    assert main(["import", str(FIX / "coincheck_youshiki1.csv"), "-o", str(ledger), "--force"]) == 0


def test_xlsx_cells_are_numbers_not_text(tmp_path):
    ledger = import_fixtures(tmp_path, "okj_orders.csv")
    xlsx = tmp_path / "out.xlsx"
    assert main(["calc", str(ledger), "--xlsx", str(xlsx)]) == 0
    book = load_workbook(xlsx)
    btc_row = [c.value for c in book["計算書_2026"][3]]
    assert btc_row[0] == "BTC"
    assert all(isinstance(v, (int, float)) for v in btc_row[1:])
    detail = [c.value for c in book["明細"][2]]
    assert isinstance(detail[6], (int, float)) and isinstance(detail[7], (int, float))


def test_same_ledger_passed_twice_is_rejected_as_duplicate_ids(tmp_path):
    import pytest
    ledger = import_fixtures(tmp_path, "okj_orders.csv")
    with pytest.raises(SystemExit, match="重複"):
        main(["calc", str(ledger), str(ledger)])
